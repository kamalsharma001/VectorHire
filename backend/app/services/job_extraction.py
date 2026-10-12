"""
VectorHire Job Description File Text Extraction Service.
Supports in-memory text extraction for PDF, DOCX, and TXT files.
Ensures conservative Unicode sanitization, error isolation, and format validation.
"""

import io
import os
from typing import Tuple
import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.exceptions import DocumentProcessingError
from app.services.document_processing import PDF_MAGIC_BYTES, sanitize_text


def extract_job_text_from_bytes(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    Extracts and sanitizes text from in-memory file bytes.
    Supported formats: .pdf, .docx, .txt.
    
    Returns:
        Tuple[str, str]: (sanitized_text, detected_format)
        
    Raises:
        DocumentProcessingError: on empty, unsupported, corrupted, or unreadable files.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise DocumentProcessingError(f"Uploaded file '{filename}' is empty (0 bytes).")

    ext = os.path.splitext(filename)[1].lower().strip()
    extracted_text = ""
    detected_format = ""

    if ext == ".txt":
        detected_format = "txt"
        try:
            extracted_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                extracted_text = file_bytes.decode("latin-1")
            except Exception as exc:
                raise DocumentProcessingError(f"Failed to decode text file '{filename}': {str(exc)}") from exc

    elif ext == ".pdf":
        detected_format = "pdf"
        # Validate magic byte
        header = file_bytes[:1024]
        if PDF_MAGIC_BYTES not in header:
            raise DocumentProcessingError(f"File '{filename}' does not have a valid PDF header signature.")

        try:
            stream = io.BytesIO(file_bytes)
            reader = PdfReader(stream)

            if reader.is_encrypted:
                try:
                    decrypted = reader.decrypt("")
                    if decrypted == 0:
                        raise DocumentProcessingError(f"PDF '{filename}' is password-protected and cannot be read.")
                except Exception as e:
                    raise DocumentProcessingError(f"Failed to decrypt password-protected PDF '{filename}': {str(e)}") from e

            page_texts = []
            for idx, page in enumerate(reader.pages, start=1):
                try:
                    p_text = page.extract_text() or ""
                    if p_text.strip():
                        page_texts.append(p_text)
                except Exception:
                    pass

            extracted_text = "\n\n".join(page_texts)

        except PdfReadError as pe:
            raise DocumentProcessingError(f"Malformed or corrupt PDF file '{filename}': {str(pe)}") from pe
        except DocumentProcessingError:
            raise
        except Exception as exc:
            raise DocumentProcessingError(f"Unexpected error parsing PDF '{filename}': {str(exc)}") from exc

    elif ext == ".docx":
        detected_format = "docx"
        try:
            stream = io.BytesIO(file_bytes)
            doc = docx.Document(stream)
            parts = []

            # Extract paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    parts.append(p_text)

            # Extract table cells
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        parts.append(" | ".join(row_cells))

            extracted_text = "\n".join(parts)

        except Exception as exc:
            raise DocumentProcessingError(f"Malformed or corrupt Word document '{filename}': {str(exc)}") from exc

    else:
        raise DocumentProcessingError(
            f"Unsupported file format '{ext}'. VectorHire accepts PDF (.pdf), Word (.docx), or plain text (.txt) files."
        )

    # Sanitize Unicode and excessive whitespace
    cleaned = sanitize_text(extracted_text)

    if len(cleaned.strip()) < 20:
        raise DocumentProcessingError(
            f"Uploaded document '{filename}' contains fewer than 20 readable characters. "
            "If this is a scanned document without an embedded text layer, OCR is required. "
            "Please paste the job description text directly."
        )

    return cleaned.strip(), detected_format
